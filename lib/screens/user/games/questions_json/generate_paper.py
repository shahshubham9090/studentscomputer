import json
import glob
import random
import os
import sys

def generate():
    questions = []
    seen_texts = set()
    
    # Locate all JSON files
    search_path = "*.json"
    files = glob.glob(search_path)
    print(f"Scanning {len(files)} JSON files in {os.getcwd()}...")
    
    for fpath in files:
        try:
            with open(fpath, 'r', encoding='utf-8') as f:
                data = json.load(f)
                if isinstance(data, list):
                    for q in data:
                        # Normalize text for deduplication
                        if isinstance(q, dict) and 'text' in q and q['text']:
                            q_text = str(q['text']).strip()
                            if q_text and q_text not in seen_texts:
                                questions.append(q)
                                seen_texts.add(q_text)
        except Exception as e:
            print(f"Error reading {fpath}: {e}")

    print(f"Total unique questions found: {len(questions)}")
    
    selected = []
    if len(questions) < 100:
        print("Warning: Fewer than 100 questions available.")
        selected = questions
    else:
        selected = random.sample(questions, 100)

    # Try to import python-docx
    try:
        from docx import Document
        from docx.shared import Pt
        
        doc = Document()
        doc.add_heading('Computer Science Question Paper', 0)
        p = doc.add_paragraph('Duration: 1 Hour')
        p.add_run('\t\t\t\t\t\tTotal Marks: 100')
        doc.add_paragraph('-' * 80)
        
        for i, q in enumerate(selected, 1):
            p = doc.add_paragraph()
            run = p.add_run(f"Q{i}. {q['text']}")
            run.bold = True
            
            choices = q.get('choices', [])
            if choices:
                formatted_choices = "   ".join([f"{chr(97+j)}) {choice}" for j, choice in enumerate(choices)])
                doc.add_paragraph(f"   {formatted_choices}")
            doc.add_paragraph() # Spacer
            
        doc.add_page_break()
        doc.add_heading('Answer Key', 0)
        
        for i, q in enumerate(selected, 1):
            correct_idx = q.get('correctChoiceIndex', -1)
            choices = q.get('choices', [])
            if isinstance(correct_idx, int) and 0 <= correct_idx < len(choices):
                ans_text = choices[correct_idx]
                key_char = chr(97+correct_idx)
            else:
                ans_text = "Unknown"
                key_char = "?"
                
            doc.add_paragraph(f"{i}. ({key_char}) {ans_text}")
            
        doc.save('question_paper.docx')
        print("Successfully generated question_paper.docx")

        # Also generate Markdown for quick preview in IDE
        with open('question_paper.md', 'w', encoding='utf-8') as f:
            f.write("# Computer Science Question Paper\n\n")
            f.write("Duration: 1 Hour | Total Marks: 100\n\n---\n\n")
            
            for i, q in enumerate(selected, 1):
                f.write(f"Q{i}. {q['text']}\n\n")
                choices = q.get('choices', [])
                if choices:
                    formatted_choices = "   ".join([f"{chr(97+j)}) {choice}" for j, choice in enumerate(choices)])
                    f.write(f"   {formatted_choices}\n")
                f.write("\n")
            
            f.write("\n---\n# Answer Key\n\n")
            for i, q in enumerate(selected, 1):
                correct_idx = q.get('correctChoiceIndex', -1)
                choices = q.get('choices', [])
                if isinstance(correct_idx, int) and 0 <= correct_idx < len(choices):
                    ans_text = choices[correct_idx]
                    key_char = chr(97+correct_idx)
                else:
                    ans_text = "Unknown"
                    key_char = "?"
                f.write(f"{i}. ({key_char}) {ans_text}\n")
        print("Successfully generated question_paper.md")
        
    except ImportError:
        print("python-docx library not found. Generating HTML-based .doc file as fallback.")
        
        # HTML disguised as DOC (Word can open this)
        # Note: We name it .doc so Word opens it by default as a rich document.
        with open('question_paper.doc', 'w', encoding='utf-8') as f:
             f.write("<html><body style='font-family:Arial, sans-serif'>")
             f.write("<h1 style='text-align:center'>Computer Science Question Paper</h1>")
             f.write("<p style='text-align:center'>Duration: 1 Hour | Total Marks: 100</p><hr/>")
             
             for i, q in enumerate(selected, 1):
                f.write(f"<p><b>Q{i}. {q['text']}</b></p>")
                choices = q.get('choices', [])
                if choices:
                    formatted_choices = "&nbsp;&nbsp;&nbsp;".join([f"{chr(97+j)}) {choice}" for j, choice in enumerate(choices)])
                    f.write(f"<div style='margin-left:20px'>{formatted_choices}</div>")
                f.write("<br/>")
             
             f.write("<br/><hr/><h1 style='text-align:center'>Answer Key</h1>")
             for i, q in enumerate(selected, 1):
                correct_idx = q.get('correctChoiceIndex', -1)
                choices = q.get('choices', [])
                if isinstance(correct_idx, int) and 0 <= correct_idx < len(choices):
                    ans_text = choices[correct_idx]
                    key_char = chr(97+correct_idx)
                else:
                    ans_text = "Unknown"
                    key_char = "?"
                f.write(f"<p>{i}. ({key_char}) {ans_text}</p>")

             f.write("</body></html>")
        print("Generated question_paper.doc (HTML format compatible with Word)")

if __name__ == "__main__":
    generate()
